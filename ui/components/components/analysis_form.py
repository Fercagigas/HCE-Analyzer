
"""
Formulario de análisis de historias clínicas
"""
import streamlit as st
import tempfile
import os
from typing import Dict, Any, Optional
from datetime import datetime
from src.analyzers.agents import AnalysisAgent
from services.auth.session_manager import SessionManager
from utils.helpers.utils import FileUtils, ValidationUtils
from config.config import ANALYSIS_TYPES, RATE_LIMITS, UI_MESSAGES

class AnalysisForm:
    """Formulario para análisis de historias clínicas"""
    
    def __init__(self):
        self.analysis_agent = AnalysisAgent()
        self.supported_formats = ['.pdf', '.txt', '.docx']
        self.max_file_size_mb = 10
    
    def render_form(self):
        """Renderiza el formulario principal de análisis"""
        st.subheader("📋 Nuevo Análisis Clínico")
        
        # Verificar límites de uso
        if not self._check_usage_limits():
            return
        
        # Tabs para diferentes tipos de entrada
        tab1, tab2 = st.tabs(["📄 Subir Archivo", "✍️ Texto Directo"])
        
        with tab1:
            self._render_file_upload_form()
        
        with tab2:
            self._render_text_input_form()
    
    def _check_usage_limits(self) -> bool:
        """Verifica los límites de uso diario"""
        if not SessionManager.is_authenticated():
            st.error(UI_MESSAGES["errors"]["auth_required"])
            return False
        
        user_id = st.session_state.user['id']
        can_analyze, remaining = self.analysis_agent.check_rate_limit(user_id)
        
        if not can_analyze:
            st.error(UI_MESSAGES["errors"]["rate_limit_exceeded"])
            st.info(f"💡 Límite diario: {RATE_LIMITS['daily_analyses']} análisis")
            return False
        
        # Mostrar análisis restantes
        st.info(f"📊 Análisis restantes hoy: **{remaining}**")
        return True
    
    def _render_file_upload_form(self):
        """Formulario para subir archivos"""
        st.markdown("### 📤 Subir Historia Clínica")
        
        # Selector de tipo de análisis
        analysis_type = st.selectbox(
            "🔬 Tipo de Análisis",
            list(ANALYSIS_TYPES.keys()),
            format_func=lambda x: f"{ANALYSIS_TYPES[x]['icon']} {ANALYSIS_TYPES[x]['name']}",
            help="Selecciona el tipo de análisis más apropiado para obtener mejores resultados"
        )
        
        # Upload de archivo
        uploaded_file = st.file_uploader(
            "Selecciona el archivo de la historia clínica",
            type=['pdf', 'txt', 'docx'],
            help=f"Formatos soportados: {', '.join(self.supported_formats)}. Tamaño máximo: {self.max_file_size_mb}MB"
        )
        
        if uploaded_file is not None:
            # Mostrar información del archivo
            self._show_file_info(uploaded_file)
            
            # Configuraciones adicionales
            with st.expander("⚙️ Configuraciones Avanzadas", expanded=False):
                include_rag = st.checkbox(
                    "🔍 Incluir consulta a guías clínicas",
                    value=True,
                    help="Enriquece el análisis consultando guías clínicas relevantes"
                )
                
                specialty_filter = st.selectbox(
                    "🏥 Filtrar por especialidad",
                    ["Automático", "Urgencias", "Cardiología", "Neurología", "Pediatría", "Otros"],
                    help="Filtra las guías clínicas por especialidad específica"
                )
                
                detail_level = st.select_slider(
                    "📊 Nivel de detalle",
                    options=["Básico", "Intermedio", "Avanzado", "Completo"],
                    value="Intermedio",
                    help="Controla la profundidad del análisis"
                )
            
            # Botón de análisis
            col1, col2, col3 = st.columns([1, 2, 1])
            with col2:
                if st.button(
                    f"🚀 Analizar {ANALYSIS_TYPES[analysis_type]['name']}",
                    use_container_width=True,
                    type="primary"
                ):
                    self._process_file_analysis(
                        uploaded_file, 
                        analysis_type, 
                        include_rag, 
                        specialty_filter,
                        detail_level
                    )
    
    def _render_text_input_form(self):
        """Formulario para entrada de texto directo"""
        st.markdown("### ✍️ Introducir Datos Manualmente")
        
        # Selector de tipo de análisis
        analysis_type = st.selectbox(
            "🔬 Tipo de Análisis",
            list(ANALYSIS_TYPES.keys()),
            format_func=lambda x: f"{ANALYSIS_TYPES[x]['icon']} {ANALYSIS_TYPES[x]['name']}",
            key="text_analysis_type"
        )
        
        # Área de texto
        clinical_data = st.text_area(
            "📝 Datos Clínicos",
            height=300,
            placeholder=self._get_placeholder_text(analysis_type),
            help="Introduce los datos clínicos que deseas analizar"
        )
        
        if clinical_data:
            # Mostrar estadísticas del texto
            self._show_text_stats(clinical_data)
            
            # Configuraciones
            with st.expander("⚙️ Configuraciones", expanded=False):
                include_rag = st.checkbox(
                    "🔍 Consultar guías clínicas",
                    value=True,
                    key="text_include_rag"
                )
                
                specialty_filter = st.selectbox(
                    "🏥 Especialidad",
                    ["Automático", "Urgencias", "Cardiología", "Neurología", "Pediatría", "Otros"],
                    key="text_specialty_filter"
                )
            
            # Botón de análisis
            col1, col2, col3 = st.columns([1, 2, 1])
            with col2:
                if st.button(
                    f"🔍 Analizar Texto",
                    use_container_width=True,
                    type="primary",
                    key="analyze_text_button"
                ):
                    self._process_text_analysis(
                        clinical_data,
                        analysis_type,
                        include_rag,
                        specialty_filter
                    )
    
    def _show_file_info(self, uploaded_file):
        """Muestra información del archivo subido"""
        file_size_mb = len(uploaded_file.getvalue()) / (1024 * 1024)
        
        col1, col2, col3 = st.columns(3)
        
        with col1:
            st.metric("📄 Archivo", uploaded_file.name)
        
        with col2:
            st.metric("📏 Tamaño", f"{file_size_mb:.2f} MB")
        
        with col3:
            st.metric("📅 Subido", datetime.now().strftime("%H:%M"))
        
        # Validar archivo
        if file_size_mb > self.max_file_size_mb:
            st.error(f"❌ Archivo muy grande. Máximo permitido: {self.max_file_size_mb}MB")
            return False
        
        return True
    
    def _show_text_stats(self, text: str):
        """Muestra estadísticas del texto introducido"""
        from utils.helpers.utils import AnalyticsUtils
        
        stats = AnalyticsUtils.calculate_text_stats(text)
        
        col1, col2, col3, col4 = st.columns(4)
        
        with col1:
            st.metric("📝 Palabras", stats['words'])
        
        with col2:
            st.metric("📄 Caracteres", stats['characters'])
        
        with col3:
            st.metric("📋 Párrafos", stats['paragraphs'])
        
        with col4:
            st.metric("⏱️ Lectura", f"{stats['reading_time_minutes']} min")
    
    def _get_placeholder_text(self, analysis_type: str) -> str:
        """Obtiene texto de placeholder según el tipo de análisis"""
        placeholders = {
            "blood_test": """Ejemplo de análisis de sangre:

Hemograma Completo:
- Hemoglobina: 12.5 g/dL (12.0-15.5)
- Hematocrito: 37.8% (36.0-46.0)
- Leucocitos: 8,500/μL (4,500-11,000)
- Plaquetas: 280,000/μL (150,000-450,000)

Bioquímica:
- Glucosa: 95 mg/dL (70-100)
- Creatinina: 0.9 mg/dL (0.6-1.2)
- Colesterol total: 185 mg/dL (<200)
- Triglicéridos: 120 mg/dL (<150)""",
            
            "imaging": """Ejemplo de reporte de imagen:

RADIOGRAFÍA DE TÓRAX PA Y LATERAL
Técnica: Radiografía digital
Calidad: Adecuada

HALLAZGOS:
- Silueta cardíaca de tamaño normal
- Campos pulmonares bien expandidos
- No se observan infiltrados ni consolidaciones
- Hilios pulmonares de aspecto normal
- Senos costofrénicos libres
- Estructuras óseas sin alteraciones

IMPRESIÓN:
Radiografía de tórax sin alteraciones patológicas evidentes""",
            
            "general_report": """Ejemplo de reporte médico:

HISTORIA CLÍNICA
Paciente: Masculino, 45 años
Motivo de consulta: Dolor torácico de 2 horas de evolución

ANTECEDENTES:
- HTA en tratamiento con enalapril
- Tabaquismo: 20 cigarrillos/día por 20 años
- Padre con IAM a los 50 años

EXAMEN FÍSICO:
- PA: 150/90 mmHg, FC: 85 lpm
- Dolor precordial opresivo 7/10
- Ruidos cardíacos rítmicos, no soplos
- Pulmones: MV conservado bilateral

ESTUDIOS:
- ECG: Ritmo sinusal, sin alteraciones agudas
- Troponinas: Pendientes""",
            
            "pathology": """Ejemplo de reporte anatomopatológico:

BIOPSIA DE PIEL
Sitio: Brazo derecho
Técnica: Punch biopsy 4mm

DESCRIPCIÓN MACROSCÓPICA:
Fragmento de piel de 4x3x2mm, superficie irregular

DESCRIPCIÓN MICROSCÓPICA:
Epidermis con hiperqueratosis y acantosis moderada.
Dermis con infiltrado inflamatorio linfocitario perivascular.
No se observa atipia celular ni signos de malignidad.

DIAGNÓSTICO:
Dermatitis crónica inespecífica"""
        }
        
        return placeholders.get(analysis_type, "Introduce aquí los datos clínicos que deseas analizar...")
    
    def _process_file_analysis(self, uploaded_file, analysis_type: str, include_rag: bool, specialty_filter: str, detail_level: str):
        """Procesa análisis de archivo"""
        try:
            # Crear archivo temporal
            with tempfile.NamedTemporaryFile(delete=False, suffix=f".{uploaded_file.name.split('.')[-1]}") as tmp_file:
                tmp_file.write(uploaded_file.getvalue())
                tmp_file_path = tmp_file.name
            
            # Extraer texto del archivo
            with st.spinner("📄 Extrayendo texto del archivo..."):
                text_content = self._extract_text_from_file(tmp_file_path, uploaded_file.type)
            
            # Limpiar archivo temporal
            os.unlink(tmp_file_path)
            
            if not text_content:
                st.error("❌ No se pudo extraer texto del archivo")
                return
            
            # Procesar análisis
            self._perform_analysis(text_content, analysis_type, include_rag, specialty_filter, detail_level)
            
        except Exception as e:
            st.error(f"❌ Error procesando archivo: {str(e)}")
    
    def _process_text_analysis(self, text: str, analysis_type: str, include_rag: bool, specialty_filter: str):
        """Procesa análisis de texto directo"""
        try:
            # Sanitizar texto
            clean_text = ValidationUtils.sanitize_input(text, max_length=10000)
            
            if not clean_text:
                st.error("❌ El texto está vacío o no es válido")
                return
            
            # Procesar análisis
            self._perform_analysis(clean_text, analysis_type, include_rag, specialty_filter, "Intermedio")
            
        except Exception as e:
            st.error(f"❌ Error procesando texto: {str(e)}")
    
    def _extract_text_from_file(self, file_path: str, file_type: str) -> str:
        """Extrae texto de diferentes tipos de archivo"""
        try:
            if file_type == "application/pdf":
                import PyPDF2
                text = ""
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page in pdf_reader.pages:
                        text += page.extract_text() + "\n"
                return text
            
            elif file_type == "text/plain":
                with open(file_path, 'r', encoding='utf-8') as file:
                    return file.read()
            
            elif file_type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document"]:
                # Para archivos Word (requiere python-docx)
                try:
                    import docx
                    doc = docx.Document(file_path)
                    text = ""
                    for paragraph in doc.paragraphs:
                        text += paragraph.text + "\n"
                    return text
                except ImportError:
                    st.error("❌ Soporte para archivos Word no disponible")
                    return ""
            
            else:
                st.error(f"❌ Tipo de archivo no soportado: {file_type}")
                return ""
                
        except Exception as e:
            st.error(f"❌ Error extrayendo texto: {str(e)}")
            return ""
    
    def _perform_analysis(self, text: str, analysis_type: str, include_rag: bool, specialty_filter: str, detail_level: str):
        """Realiza el análisis clínico"""
        user_id = st.session_state.user['id']
        session_id = st.session_state.current_session['id']
        
        # Mostrar progreso
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        try:
            # Paso 1: Preparar análisis
            status_text.text("🔄 Preparando análisis...")
            progress_bar.progress(20)
            
            # Paso 2: Análisis principal
            status_text.text("🧠 Analizando con IA...")
            progress_bar.progress(60)
            
            result = self.analysis_agent.analyze_clinical_data(
                user_id=user_id,
                analysis_type=analysis_type,
                data=text
            )
            
            # Paso 3: Enriquecer con RAG si está habilitado
            if include_rag and result.get('success'):
                status_text.text("📚 Consultando guías clínicas...")
                progress_bar.progress(80)
                
                # Aquí se integraría la consulta RAG
                # rag_context = self._get_rag_context(text, specialty_filter)
                # result = self._enrich_with_rag(result, rag_context)
            
            # Paso 4: Finalizar
            status_text.text("✅ Análisis completado")
            progress_bar.progress(100)
            
            # Mostrar resultados
            self._display_analysis_results(result, analysis_type)
            
            # Guardar en sesión
            if result.get('success'):
                self._save_analysis_to_session(session_id, text, result, analysis_type)
            
        except Exception as e:
            st.error(f"❌ Error en análisis: {str(e)}")
        finally:
            # Limpiar elementos de progreso
            progress_bar.empty()
            status_text.empty()
    
    def _display_analysis_results(self, result: Dict[str, Any], analysis_type: str):
        """Muestra los resultados del análisis"""
        if not result.get('success'):
            st.error(f"❌ Error en análisis: {result.get('error', 'Error desconocido')}")
            return
        
        # Título del resultado
        st.success("✅ Análisis completado exitosamente")
        
        # Información del análisis
        col1, col2, col3 = st.columns(3)
        
        with col1:
            st.metric("🤖 Modelo", result.get('model_used', 'Desconocido'))
        
        with col2:
            st.metric("📊 Tipo", ANALYSIS_TYPES[analysis_type]['name'])
        
        with col3:
            remaining = result.get('remaining_analyses', 0)
            st.metric("🔄 Restantes", remaining)
        
        # Resultado principal
        st.markdown("### 📋 Resultado del Análisis")
        
        # Formatear y mostrar el análisis
        analysis_text = result.get('analysis', '')
        
        # Dividir en secciones si tiene formato estructurado
        if "**" in analysis_text:
            self._display_structured_analysis(analysis_text)
        else:
            st.markdown(analysis_text)
        
        # Botones de acción
        col1, col2, col3 = st.columns(3)
        
        with col1:
            if st.button("💾 Guardar Análisis", use_container_width=True):
                st.success("✅ Análisis guardado en tu historial")
        
        with col2:
            if st.button("📤 Exportar PDF", use_container_width=True):
                # Implementar exportación a PDF
                st.info("🔄 Funcionalidad en desarrollo")
        
        with col3:
            if st.button("🔄 Nuevo Análisis", use_container_width=True):
                st.rerun()
    
    def _display_structured_analysis(self, analysis_text: str):
        """Muestra análisis estructurado con formato mejorado"""
        sections = analysis_text.split("**")
        
        current_section = ""
        for i, section in enumerate(sections):
            if i % 2 == 0:  # Texto normal
                if section.strip():
                    st.markdown(section.strip())
            else:  # Encabezado
                if section.strip():
                    st.markdown(f"### {section.strip()}")
    
    def _save_analysis_to_session(self, session_id: str, input_data: str, result: Dict[str, Any], analysis_type: str):
        """Guarda el análisis en la sesión actual"""
        try:
            # Guardar mensaje del usuario
            st.session_state.auth_service.save_message(
                session_id, 
                f"Análisis de {ANALYSIS_TYPES[analysis_type]['name']}: {input_data[:100]}...", 
                'user'
            )
            
            # Guardar respuesta del asistente
            st.session_state.auth_service.save_message(
                session_id,
                result.get('analysis', ''),
                'assistant'
            )
            
        except Exception as e:
            st.warning(f"⚠️ No se pudo guardar en historial: {str(e)}")

def show_analysis_form():
    """Función principal para mostrar el análisis con interfaz de chat mejorada"""
    try:
        # Verificar autenticación
        if not SessionManager.is_authenticated():
            st.error("❌ Debes iniciar sesión para realizar análisis")
            return
        
        # Verificar sesión activa
        if not st.session_state.get('current_session'):
            st.error("❌ No hay sesión activa. Crea una nueva sesión desde el inicio.")
            return
        
        # Inicializar el historial de análisis en session_state si no existe
        if "analysis_messages" not in st.session_state:
            st.session_state.analysis_messages = []
        
        # Sidebar con opciones de análisis
        with st.sidebar:
            st.subheader("⚙️ Configuración de Análisis")
            
            analysis_type = st.selectbox(
                "Tipo de Análisis",
                list(ANALYSIS_TYPES.keys()),
                format_func=lambda x: f"{ANALYSIS_TYPES[x]['icon']} {ANALYSIS_TYPES[x]['name']}",
                help="Selecciona el tipo de análisis médico"
            )
            
            include_rag = st.checkbox(
                "🔍 Consultar guías clínicas",
                value=True,
                help="Enriquece el análisis con guías clínicas"
            )
            
            st.divider()
            
            # Botón para limpiar historial
            if st.button("🗑️ Limpiar Historial", use_container_width=True):
                st.session_state.analysis_messages = []
                st.rerun()
            
            # Estadísticas
            if st.session_state.analysis_messages:
                analyses_count = len([m for m in st.session_state.analysis_messages if m["role"] == "assistant"])
                st.metric("📊 Análisis realizados", analyses_count)
        
        # Cargar historial de la base de datos
        _load_analysis_history()
        
        # Mostrar historial de análisis
        for message in st.session_state.analysis_messages:
            with st.chat_message(message["role"]):
                if message["role"] == "assistant":
                    _render_analysis_result(message["content"])
                else:
                    st.write(message["content"])
        
        # Opciones de entrada si no hay mensajes
        if not st.session_state.analysis_messages:
            st.markdown("### 📋 Opciones de Análisis")
            
            col1, col2 = st.columns(2)
            
            with col1:
                st.markdown("**📄 Subir Archivo**")
                uploaded_file = st.file_uploader(
                    "Selecciona un archivo médico",
                    type=['pdf', 'txt', 'docx'],
                    help="Formatos soportados: PDF, TXT, DOCX"
                )
                
                if uploaded_file and st.button("🔍 Analizar Archivo", use_container_width=True):
                    _process_file_analysis(uploaded_file, analysis_type, include_rag)
            
            with col2:
                st.markdown("**✍️ Texto Directo**")
                text_input = st.text_area(
                    "Pega el texto médico aquí",
                    height=150,
                    placeholder="Ejemplo: Paciente de 45 años con dolor torácico..."
                )
                
                if text_input and st.button("🔍 Analizar Texto", use_container_width=True):
                    _process_text_analysis(text_input, analysis_type, include_rag)
        
        # Input de chat para preguntas adicionales
        if st.session_state.analysis_messages:
            if prompt := st.chat_input("Haz una pregunta sobre el análisis o sube un nuevo documento..."):
                # Agregar pregunta al historial
                st.session_state.analysis_messages.append({"role": "user", "content": prompt})
                
                # Mostrar pregunta
                with st.chat_message("user"):
                    st.write(prompt)
                
                # Procesar pregunta de seguimiento
                _process_followup_question(prompt, analysis_type)
        
    except Exception as e:
        st.error(f"❌ Error inicializando análisis: {str(e)}")
        if st.checkbox("Mostrar detalles del error"):
            st.exception(e)

def _load_analysis_history():
    """Carga el historial de análisis de la base de datos"""
    if not st.session_state.get('current_session'):
        return
    
    # Solo cargar si el historial local está vacío
    if not st.session_state.analysis_messages:
        try:
            success, messages = st.session_state.auth_service.get_session_messages(
                st.session_state.current_session['id']
            )
            
            if success and messages:
                # Convertir mensajes de BD al formato del chat
                for message in messages:
                    st.session_state.analysis_messages.append({
                        "role": message['role'],
                        "content": message['content']
                    })
        except Exception as e:
            st.warning(f"⚠️ No se pudo cargar el historial: {str(e)}")

def _process_file_analysis(uploaded_file, analysis_type: str, include_rag: bool):
    """Procesa el análisis de un archivo"""
    try:
        # Agregar mensaje del usuario
        file_message = f"📄 Archivo subido: **{uploaded_file.name}** ({ANALYSIS_TYPES[analysis_type]['name']})"
        st.session_state.analysis_messages.append({"role": "user", "content": file_message})
        
        # Mostrar mensaje del usuario
        with st.chat_message("user"):
            st.write(file_message)
        
        # Procesar análisis
        with st.chat_message("assistant"):
            with st.spinner("🔍 Analizando archivo médico..."):
                # Simular análisis (aquí iría la lógica real)
                import time
                time.sleep(2)  # Simular procesamiento
                
                result = {
                    "success": True,
                    "analysis": f"""### 📋 Análisis de {ANALYSIS_TYPES[analysis_type]['name']}

**Archivo procesado:** {uploaded_file.name}

**Hallazgos principales:**
• Documento médico válido procesado correctamente
• Estructura de datos clínicos identificada
• Información relevante extraída para análisis

**Recomendaciones:**
• Revisar los hallazgos con el contexto clínico completo
• Considerar estudios complementarios si es necesario
• Seguimiento según protocolos establecidos

**Nota:** Este es un análisis simulado. En la versión completa se procesaría el contenido real del archivo.""",
                    "type": analysis_type,
                    "timestamp": datetime.now().isoformat()
                }
                
                # Agregar resultado al historial
                st.session_state.analysis_messages.append({
                    "role": "assistant", 
                    "content": result
                })
                
                # Mostrar resultado
                _render_analysis_result(result)
                
                # Guardar en base de datos
                _save_to_database(file_message, result)
        
    except Exception as e:
        st.error(f"Error procesando archivo: {str(e)}")

def _process_text_analysis(text: str, analysis_type: str, include_rag: bool):
    """Procesa el análisis de texto directo"""
    try:
        # Agregar mensaje del usuario
        user_message = f"✍️ **Texto para análisis** ({ANALYSIS_TYPES[analysis_type]['name']}):\n\n{text[:200]}{'...' if len(text) > 200 else ''}"
        st.session_state.analysis_messages.append({"role": "user", "content": user_message})
        
        # Mostrar mensaje del usuario
        with st.chat_message("user"):
            st.write(user_message)
        
        # Procesar análisis
        with st.chat_message("assistant"):
            with st.spinner("🔍 Analizando texto médico..."):
                # Simular análisis
                import time
                time.sleep(2)
                
                result = {
                    "success": True,
                    "analysis": f"""### 📋 Análisis de {ANALYSIS_TYPES[analysis_type]['name']}

**Texto analizado:** {len(text)} caracteres

**Interpretación clínica:**
• Datos clínicos procesados correctamente
• Información médica relevante identificada
• Contexto clínico evaluado

**Hallazgos destacados:**
• Parámetros dentro de rangos esperados
• Correlación clínica apropiada
• Seguimiento recomendado según protocolo

**Recomendaciones:**
• Continuar monitoreo según indicaciones
• Evaluar necesidad de estudios adicionales
• Seguimiento clínico programado

**Nota:** Este es un análisis simulado para demostración.""",
                    "type": analysis_type,
                    "timestamp": datetime.now().isoformat()
                }
                
                # Agregar resultado al historial
                st.session_state.analysis_messages.append({
                    "role": "assistant", 
                    "content": result
                })
                
                # Mostrar resultado
                _render_analysis_result(result)
                
                # Guardar en base de datos
                _save_to_database(user_message, result)
        
    except Exception as e:
        st.error(f"Error procesando texto: {str(e)}")

def _process_followup_question(question: str, analysis_type: str):
    """Procesa preguntas de seguimiento sobre el análisis"""
    try:
        with st.chat_message("assistant"):
            with st.spinner("🤔 Procesando pregunta..."):
                # Simular respuesta inteligente
                import time
                time.sleep(1)
                
                response = f"""Respuesta a tu consulta: "{question}"

Basándome en el análisis previo de {ANALYSIS_TYPES[analysis_type]['name']}, puedo proporcionarte la siguiente información:

• La pregunta está relacionada con el contexto clínico analizado
• Los hallazgos previos son relevantes para esta consulta
• Se recomienda considerar el contexto completo del paciente

¿Hay algún aspecto específico que te gustaría que profundice más?"""
                
                # Agregar respuesta al historial
                st.session_state.analysis_messages.append({
                    "role": "assistant", 
                    "content": response
                })
                
                st.write(response)
                
                # Guardar en base de datos
                _save_to_database(question, response)
        
    except Exception as e:
        st.error(f"Error procesando pregunta: {str(e)}")

def _render_analysis_result(content):
    """Renderiza el resultado del análisis"""
    if isinstance(content, dict):
        # Resultado estructurado
        if content.get("success"):
            st.success("✅ Análisis completado exitosamente")
            
            if content.get("analysis"):
                st.markdown(content["analysis"])
            
            # Mostrar metadata si está disponible
            if content.get("timestamp"):
                st.caption(f"🕒 Procesado: {content['timestamp'][:19].replace('T', ' ')}")
        else:
            st.error(f"❌ Error en el análisis: {content.get('error', 'Error desconocido')}")
    else:
        # Respuesta de texto simple
        st.write(content)

def _save_to_database(user_message: str, assistant_response):
    """Guarda los mensajes en la base de datos"""
    try:
        if st.session_state.get('current_session'):
            session_id = st.session_state.current_session['id']
            
            # Guardar mensaje del usuario
            st.session_state.auth_service.save_message(session_id, user_message, 'user')
            
            # Guardar respuesta del asistente
            response_text = assistant_response if isinstance(assistant_response, str) else assistant_response.get('analysis', str(assistant_response))
            st.session_state.auth_service.save_message(session_id, response_text, 'assistant')
            
    except Exception as e:
        st.warning(f"⚠️ No se pudo guardar en la base de datos: {str(e)}")

